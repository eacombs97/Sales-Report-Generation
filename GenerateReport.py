import os
import pandas as pd
import logging
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from datetime import datetime

###################
#FIELDS & CONFIGURATION
####################
INPUT_DIR = #path to Sales report CV
OUTPUT_DIR = #path to outputted report
REPORT_NAME = #Report title
FILENAME_DATE_FORMAT = "%Y-%m-%d"

# Logging Configuration
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s",
    handlers=[
        logging.FileHandler("report_tool.log"),
        logging.StreamHandler()
    ]
)

###################
#MAKE SALES REPORT METHODS
####################
#Make sure directory exists and if it doesn't, create one 
os.makedirs(OUTPUT_DIR, exist_ok=True)

#Load data from CV
def load_data(directory):
    """Load and concatenate all CSV files from the directory."""
    all_data = []
    for file in os.listdir(directory):
        if file.endswith(".csv"):
            df = pd.read_csv(os.path.join(directory, file))
            all_data.append(df)
    return pd.concat(all_data, ignore_index=True)
    
def transform_data(df):
    """Apply business logic to transform raw data."""
    df["Total"] = df["Quantity"] * df["Unit_Price"]
    df["Date"] = pd.to_datetime(df["Date"])
    df = df.sort_values(by="Date")
    return df
    
#Create Report
def write_excel_report(df, filename):
    """Write the DataFrame to a styled Excel file."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales Summary"

    # Write headers
    headers = list(df.columns)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", fill_type="solid")

    ws.append(headers)
    for col_num, cell in enumerate(ws[1], 1):
        cell.font = header_font
        cell.fill = header_fill
        ws.column_dimensions[cell.column_letter].width = 15

    # Write rows
    for row in df.itertuples(index=False):
        ws.append(row)

    # Save
    output_path = os.path.join(OUTPUT_DIR, filename)
    wb.save(output_path)
    print(f"[✓] Report saved: {output_path}")
    logging.info(f"[✓] Report saved: {output_path}")
    # Add report information to logging

###################
#CREATE POWERPOINT
####################
def generate_ppt_summary(df, filename):
    prs = Presentation()
    
    # --- Slide 1: Title Slide ---
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Monthly Sales Report"
    slide.placeholders[1].text = f"Generated on {datetime.today().strftime('%B %d, %Y')}"

    # --- Slide 2: Bar Chart (Sales by Region) ---
    slide_layout = prs.slide_layouts[5]
    slide = prs.slides.add_slide(slide_layout)
    shapes = slide.shapes
    shapes.title.text = "Total Sales by Region"

    sales_by_region = df.groupby("Region")["Total"].sum()

    chart_data = CategoryChartData()
    chart_data.categories = list(sales_by_region.index)
    chart_data.add_series("Sales", list(sales_by_region.values))

    x, y, cx, cy = Inches(1), Inches(1.5), Inches(8), Inches(4.5)
    chart = shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, cx, cy, chart_data
    ).chart
    chart.has_legend = False

    # --- Slide 3: Data Table (Top 10 Orders) ---
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Top 10 Orders by Value"
    
    top_orders = df.sort_values("Total", ascending=False).head(10)
    rows, cols = top_orders.shape
    x, y, cx, cy = Inches(0.5), Inches(1.5), Inches(9), Inches(4)

    table = slide.shapes.add_table(rows+1, cols, x, y, cx, cy).table

    # Header row
    for col_idx, col_name in enumerate(top_orders.columns):
        cell = table.cell(0, col_idx)
        cell.text = col_name
        cell.text_frame.paragraphs[0].font.bold = True

    # Data rows
    for i, row in top_orders.iterrows():
        for j, val in enumerate(row):
            table.cell(i+1, j).text = str(val)

    # --- Save ---
    pptx_path = os.path.join(OUTPUT_DIR, filename)
    prs.save(pptx_path)
    print(f"[✓] PowerPoint saved: {pptx_path}")
    logging.info(f"[✓] PowerPoint saved: {pptx_path}")
    
###################
#MAIN
####################
def main():
    print("[*] Loading data...")
    logging.info("[*] Loading data...")
    df = load_data(DATA_DIR)

    print("[*] Transforming data...")
    logging.info("[*] Transforming data...")
    df = transform_data(df)

    today_str = datetime.today().strftime(FILENAME_DATE_FORMAT)
    output_filename = f"{REPORT_TEMPLATE_NAME}_{today_str}.xlsx"

    print("[*] Writing Excel report...")
    logging.info("[*] Writing Excel report...")
    write_excel_report(df, output_filename)
    
    print("[*] Writing PowerPoint presentation...")
    logging.info("[*] Writing PowerPoint presentation...")
    ppt_filename = f"{REPORT_TEMPLATE_NAME}_{today_str}.pptx"
    generate_ppt_summary(df, ppt_filename)

if __name__ == "__main__":
    main()